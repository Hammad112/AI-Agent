import os
try:
    from docx import Document
    from docx.shared import Inches
except ImportError:
    print("python-docx not installed. Please install it using 'pip install python-docx'")
    exit(1)

def create_documentation():
    doc = Document()
    doc.add_heading('Generic AI Agent for Business Orders & Scheduling', 0)

    doc.add_heading('1. Project Overview', level=1)
    doc.add_paragraph(
        "This project is a modular, RAG-based AI agent designed to handle business orders, "
        "appointment scheduling, and general info Q&A based on a single PDF knowledge source."
    )

    doc.add_heading('2. Tech Stack', level=1)
    doc.add_paragraph("Python 3.11+, LangChain, LangGraph, FastAPI, SQLite, OpenAI/Gemini API.", style='List Bullet')

    doc.add_heading('3. Key Features', level=1)
    doc.add_paragraph("PDF Knowledge Enrichment: Generates supplementary 'skills' and product comparisons.", style='List Bullet')
    doc.add_paragraph("No-Embedding RAG: Uses LLM-based topic matching for high-precision retrieval.", style='List Bullet')
    doc.add_paragraph("Agentic Pipeline: Structured as Planner -> Router -> Tools -> Critic/Evaluator.", style='List Bullet')
    doc.add_paragraph("Real-world Validation: Integrated geopy for Canadian address and postal code verification.", style='List Bullet')
    doc.add_paragraph("Multi-turn Memory: Persistent conversation state and slot-filling logic.", style='List Bullet')
    doc.add_paragraph("Local Calendar: Automatic .ics file generation for all appointments.", style='List Bullet')
    doc.add_paragraph("Extensive Auditing: All LLM prompts/responses and tool calls are saved to SQLite.", style='List Bullet')

    doc.add_heading('4. Installation & Setup', level=1)
    doc.add_paragraph("1. Clone the repository.", style='List Number')
    doc.add_paragraph("2. Install dependencies: pip install -r requirements.txt", style='List Number')
    doc.add_paragraph("3. Configure .env file with your API keys.", style='List Number')
    doc.add_paragraph("4. Run the agent: python main.py --pdf <your_file.pdf>", style='List Number')

    doc.add_heading('5. Database & Logs', level=1)
    doc.add_paragraph(
        "All data is stored in business_agent.db. Logs include conversations, tool calls, "
        "agent events, and full LLM prompts/responses."
    )

    filename = 'Project_Documentation.docx'
    doc.save(filename)
    print(f"Documentation saved as {filename}")

if __name__ == "__main__":
    create_documentation()
